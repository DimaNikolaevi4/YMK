# -*- coding: utf-8 -*-
"""Доводка Т1/Т2 шести КТП (02.01/02.02/03.02 С-21/С-22) до эталонного формата 05.02.

Команда владельца «делай, редактируй все (полностью все из нагрузки) ктп» — доводка
формата после заполнения кол.9 (ищи-3). Устраняет последние 3 провала валидатора
(12а, 12в-1, 12в-2), существовавшие с момента генерации (1556486):

1. Т1: строка «Компл. дифф. зачет» — часы (2) в графе 6, остальные графы пусты.
   Реализация: трансплантация готовой строки r9 из эталона 05.02 (гарантия
   идентичных шрифтов/выравнивания; сетки tblGrid проверены на совпадение).
2. Т1: trHeight — r1=219, r2=234, r3=1758, r4=237, r5..r11=215 (r0 без высоты).
3. Т2: trHeight — r0=20, r1=230, r2=1695, остальные строки =20 (как в эталоне).

hRule не задаётся (как в эталоне) → atLeast, контент строк не обрезается.
"""
import copy
import sys
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from lxml import etree

ETALON = '/home/z/my-project/YMK/KTP/КТП МДК 05.02 М2 курс.docx'
FILES = [
    '/home/z/my-project/YMK/KTP/КТП МДК 02.01 С-21.docx',
    '/home/z/my-project/YMK/KTP/КТП МДК 02.01 С-22.docx',
    '/home/z/my-project/YMK/KTP/КТП МДК 02.02 С-21.docx',
    '/home/z/my-project/YMK/KTP/КТП МДК 02.02 С-22.docx',
    '/home/z/my-project/YMK/KTP/КТП МДК 03.02 С-21.docx',
    '/home/z/my-project/YMK/KTP/КТП МДК 03.02 С-22.docx',
]

# trHeight эталона: Т1 — индекс → высота (r0 без высоты); Т2 — первые 3 строки
T1_HEIGHTS = {1: 219, 2: 234, 3: 1758, 4: 237}
T1_DATA_H = 215          # r5..r11
T2_HEAD_HEIGHTS = [20, 230, 1695]
T2_DATA_H = 20

KDZ_ROW_SRC = 9          # r9 эталона Т1 — «Компл. дифф. зачет»


def set_trheight(tr, val):
    """Задать w:trHeight строке, соблюдая порядок детей trPr (trHeight < tblHeader)."""
    trPr = tr.find(qn('w:trPr'))
    if trPr is None:
        trPr = tr.makeelement(qn('w:trPr'), {})
        tr.insert(0, trPr)
    th = trPr.find(qn('w:trHeight'))
    if th is None:
        th = trPr.makeelement(qn('w:trHeight'), {})
        tbl_header = trPr.find(qn('w:tblHeader'))
        if tbl_header is not None:
            tbl_header.addprevious(th)
        else:
            trPr.append(th)
    th.set(qn('w:val'), str(val))


def find_t1(doc):
    for t in doc.tables:
        if t.cell(0, 0).text.strip().startswith('Междисциплинарный курс'):
            return t
    raise RuntimeError('Т1 не найдена')


def find_t2(doc):
    for t in doc.tables:
        if '№ занятия' in t.rows[0].cells[0].text:
            return t
    for t in doc.tables:
        if '№ занятия' in ' '.join(c.text for c in t.rows[0].cells):
            return t
    raise RuntimeError('Т2 не найдена')


# ── Источник: строка «Компл. дифф. зачет» эталона ──────────────────────────
et_doc = Document(ETALON)
et_t1 = find_t1(et_doc)
et_kdz_tr = et_t1.rows[KDZ_ROW_SRC].cells[0].text.strip()
assert et_kdz_tr.lower().startswith('компл'), f'r{KDZ_ROW_SRC} эталона: {et_kdz_tr!r}'
assert et_t1.rows[KDZ_ROW_SRC].cells[5].text.strip() == '2', 'часы КДЗ эталона ≠ 2'
et_xml = etree.tostring(et_t1.rows[KDZ_ROW_SRC]._tr)

for path in FILES:
    doc = Document(path)
    t1, t2 = find_t1(doc), find_t2(doc)
    short = path.split('/')[-1]

    # ── Проверки структуры перед правкой ──
    assert len(t1.rows) == 12, f'{short}: Т1 {len(t1.rows)} строк'
    assert t1.rows[10].cells[0].text.strip() == 'Практика', short
    assert t1.rows[11].cells[0].text.strip() == 'Всего', short
    assert t1.rows[5].cells[0].text.strip().startswith('МДК'), short
    for ri in range(6, 10):
        row_txt = ''.join(c.text.strip() for c in t1.rows[ri].cells)
        assert row_txt == '', f'{short}: r{ri} не пуста: {row_txt!r}'

    # ── 1. Т1: трансплантация строки «Компл. дифф. зачет» в r9 (последняя
    #      пустая строка перед «Практика» — та же позиция, что в эталоне) ──
    old_tr = t1.rows[9]._tr
    new_tr = parse_xml(et_xml)
    old_tr.addprevious(new_tr)
    t1._tbl.remove(old_tr)

    # ── 2. Т1: trHeight ──
    for ri, h in T1_HEIGHTS.items():
        set_trheight(t1.rows[ri]._tr, h)
    for ri in range(5, 12):
        set_trheight(t1.rows[ri]._tr, T1_DATA_H)

    # ── 3. Т2: trHeight (шапка 20/230/1695, данные 20) ──
    assert len(t2.rows) >= 6, f'{short}: Т2 {len(t2.rows)} строк'
    for ri, h in enumerate(T2_HEAD_HEIGHTS):
        set_trheight(t2.rows[ri]._tr, h)
    for ri in range(len(T2_HEAD_HEIGHTS), len(t2.rows)):
        set_trheight(t2.rows[ri]._tr, T2_DATA_H)

    doc.save(path)

    # ── Контроль после правки ──
    chk = Document(path)
    ct1, ct2 = find_t1(chk), find_t2(chk)
    att = [ct1.rows[9].cells[c].text.strip() for c in range(12)]
    assert att[0].lower().startswith('компл') and att[5] == '2', f'{short}: r9 {att}'
    assert all(att[c] == '' for c in (1, 2, 3, 4, 6, 7, 8, 9, 10, 11)), f'{short}: r9 лишний текст'
    hs1 = [tr.find(qn('w:trPr') + '/' + qn('w:trHeight')) for tr in ct1._tbl.findall(qn('w:tr'))]
    hs1 = [th.get(qn('w:val')) if th is not None else None for th in hs1]
    assert hs1[1:5] == ['219', '234', '1758', '237'] and all(v == '215' for v in hs1[5:]), f'{short}: Т1 h={hs1}'
    hs2 = [tr.find(qn('w:trPr') + '/' + qn('w:trHeight')) for tr in ct2._tbl.findall(qn('w:tr'))]
    hs2 = [th.get(qn('w:val')) if th is not None else None for th in hs2]
    assert hs2[:3] == ['20', '230', '1695'] and all(v is not None for v in hs2), f'{short}: Т2 h={hs2[:6]}'
    print(f'✅ {short}: r9=«Компл. дифф. зачет»(2), Т1 h={hs1[:5]}…215, Т2 h ок ({len(hs2)} строк)')

print('\nГотово: 6 файлов приведены к эталонному формату Т1/Т2.')
