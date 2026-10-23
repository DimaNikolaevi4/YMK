#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Извлечение данных эталонного КТП в JSON (формат generate_ktp_docx.py)."""
import json
import re
from docx import Document

SRC = "/home/z/my-project/YMK/KTP/КТП МДК 05.02 М 2 курс.docx"  # Обновлено 2026-10-23 (было 'М2 курс.docx' без пробела)
OUT = "/home/z/my-project/YMK/scripts/data_05_02_example.json"

doc = Document(SRC)
paras = [p.text for p in doc.paragraphs]
t1, t2 = doc.tables[1], doc.tables[2]

d = {}

# --- Титульник ---
for p in paras:
    t = p.strip()
    m = re.match(r'на __(.+?)__ семестр\(ы\)\s*_(\d+)___\s*-\s*__(\d+)___.*?_____(\d+)_+ курс', t)
    if m:
        d['semesters'], d['year_start'], d['year_end'], d['course'] = m.group(1), m.group(2), m.group(3), m.group(4)
    m = re.match(r'учебной группы \(учебных групп\)\s*_{2,}\s*(.+?)\s*_{2,}', t)
    if m:
        d['group'] = m.group(1).strip()
    m = re.match(r'Профессиональный модуль:\s*(ПМ\.(\d+)\s*.+)', t)
    if m:
        d['pm_code'] = m.group(2)
        d['pm_name'] = re.sub(r'^ПМ\.\d+\s*', '', m.group(1)).strip()
    m = re.match(r'Междисциплинарные курсы:\s*(МДК\s*([\d.]+)\s*.+)', t)
    if m:
        d['mdk_code'] = m.group(2)
        d['mdk_name'] = re.sub(r'^МДК\s*[\d.]+\s*', '', m.group(1)).strip()
    m = re.match(r'по профессии:\s*(.+)', t)
    if m:
        d['profession'] = m.group(1).strip()
    m = re.search(r'Объем образовательной программы:\s*_*(\d+)', t)
    if m:
        d['total_hours'] = int(m.group(1))
    m = re.search(r'в том числе в форме практической подготовки\s*_*(\d+)', t)
    if m:
        d['practice_hours'] = int(m.group(1))
    m = re.search(r'теоретическое обучение\s*_*(\d+)', t)
    if m:
        d['theory_hours'] = int(m.group(1))
    m = re.search(r'промежуточная аттестация в форме\s*_*\s*(.+?)\s*_*(\d+)\s*__\(часов\)', t)
    if m:
        d['attestation'] = {'form': m.group(1).strip(), 'hours': int(m.group(2))}
    m = re.search(r'Рассмотрен на заседании цикловой комиссии\s+(\S+)\s+дисциплин', t)
    if m:
        d['commission'] = m.group(1)

# --- Таблица 1: строки семестров ---
sems = []
for ri in range(5, len(t1.rows)):
    c0 = t1.cell(ri, 0).text.strip()
    if c0 in ('', 'Практика', 'Всего'):
        break
    sems.append({
        'course': t1.cell(ri, 1).text.strip(),
        'semester': t1.cell(ri, 2).text.strip(),
        'total': int(t1.cell(ri, 3).text.strip()),
        'with_teacher': int(t1.cell(ri, 4).text.strip()),
        'theory': int(t1.cell(ri, 5).text.strip()),
        'lab': t1.cell(ri, 6).text.strip(),
        'practice': int(t1.cell(ri, 7).text.strip()),
        'coursework': t1.cell(ri, 8).text.strip(),
        'self': t1.cell(ri, 9).text.strip(),
        'uch_practice': t1.cell(ri, 10).text.strip(),
        'proizv_practice': t1.cell(ri, 11).text.strip(),
    })
d['table1_semesters'] = sems

# --- Таблица 2: темы, занятия, диф.зачёт, преподаватель ---
topics = []
cur_topic = None
lesson_num = 0
teacher = ''
att_row_name = ''
for ri in range(4, len(t2.rows)):
    c0 = t2.cell(ri, 0).text.strip()
    c1 = t2.cell(ri, 1).text.strip()
    if c1.startswith('Итого'):
        break
    if c1.startswith('МДК'):
        teacher = t2.cell(ri, 10).text.strip()
        continue
    if not c0:  # заголовочная строка темы (жирная, без номера)
        cur_topic = {
            'name': c1,
            'total_hours': int(t2.cell(ri, 2).text.strip()),
            'practice_hours': int(t2.cell(ri, 3).text.strip() or 0),
            'lessons': [],
        }
        topics.append(cur_topic)
        continue
    # строка занятия или диф.зачёт
    hours = int(t2.cell(ri, 2).text.strip() or 0)
    prep = t2.cell(ri, 3).text.strip()
    ltype = t2.cell(ri, 4).text.strip()
    if not ltype and not prep:  # строка аттестации
        att_row_name = c1
        continue
    lesson_num = int(c0)
    cur_topic['lessons'].append({
        'num': lesson_num,
        'name': c1,
        'hours': hours,
        'practice_prep': int(prep) if prep else 0,
        'type': ltype,
        'ok': t2.cell(ri, 5).text.strip(),
        'pk': t2.cell(ri, 6).text.strip(),
        'equipment': t2.cell(ri, 7).text.strip(),
        'task': t2.cell(ri, 8).text.strip(),
        'control': t2.cell(ri, 9).text.strip(),
    })
if att_row_name:
    d.setdefault('attestation', {})['row_name'] = att_row_name
d['teacher_name'] = teacher
d['topics'] = topics

# --- Таблица 2а: оборудование ---
eq = []
for ri in range(2, len(doc.tables[3].rows)):
    name = doc.tables[3].cell(ri, 1).text.strip()
    if not name:
        break
    eq.append(name)
d['equipment'] = eq

# --- Литература ---
def read_lit(tbl, with_url=False):
    out = []
    for ri in range(1, len(tbl.rows)):
        code = tbl.cell(ri, 0).text.strip()
        if not code:
            continue
        item = {
            'code': code,
            'name': tbl.cell(ri, 1).text.strip(),
            'author': tbl.cell(ri, 2).text.strip(),
            'publisher': tbl.cell(ri, 3).text.strip(),
        }
        if with_url:
            item['url'] = tbl.cell(ri, 4).text.strip()
        out.append(item)
    return out

d['sources_basic'] = read_lit(doc.tables[4])
d['sources_additional'] = read_lit(doc.tables[5])
d['sources_electronic'] = read_lit(doc.tables[6], with_url=True)

with open(OUT, 'w', encoding='utf-8') as f:
    json.dump(d, f, ensure_ascii=False, indent=2)

# Контроль
n_lessons = sum(len(t['lessons']) for t in topics)
sum_h = sum(l['hours'] for t in topics for l in t['lessons'])
sum_p = sum(l['practice_prep'] for t in topics for l in t['lessons'])
print(f"MDK: {d.get('mdk_code')} | группа {d.get('group')} | семестры {d.get('semesters')} | {d.get('year_start')}-{d.get('year_end')}")
print(f"Часы: всего={d.get('total_hours')} теор={d.get('theory_hours')} практ={d.get('practice_hours')} атт={d.get('attestation')}")
print(f"Темы: {len(topics)} | занятий: {n_lessons} | Σчасов={sum_h} | Σпракт={sum_p}")
print(f"Оборудование: {len(eq)} | ОИ: {len(d['sources_basic'])} | ДИ: {len(d['sources_additional'])} | ЭИ: {len(d['sources_electronic'])}")
print(f"Сохранено: {OUT}")
